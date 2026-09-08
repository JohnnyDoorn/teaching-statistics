# -*- coding: utf-8 -*-
"""Make the shared conclusion box neutral: colour now carries meaning on this
slide (purple = one machine, green = the other), so the box that talks about
BOTH must not borrow either colour."""
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

D=sys.argv[1]
prs=Presentation(D)
GREY=RGBColor(0x8B,0x8B,0x8B)     # the grey already used for slide numbers
FILL=RGBColor(0xF2,0xF2,0xF2)     # theme lt2, neutral against the white plot

hits=0
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        if "not matter for your wallet" not in sh.text_frame.text: continue
        sh.fill.solid(); sh.fill.fore_color.rgb = FILL
        sh.line.color.rgb = GREY; sh.line.width = Pt(1.5)
        for pa in sh.text_frame.paragraphs:
            for r in pa.runs:
                r.font.size = Pt(18)          # matches the two median boxes' lead size
                r.font.name = None            # inherit PT Sans from the theme
                r.font.color.rgb = RGBColor(0,0,0)
        hits+=1
        print(f"  slide {i}: border 77BC65 -> 8B8B8B, fill F7D1D5 -> F2F2F2, text 19pt -> 18pt")
assert hits==1, f"expected 1 conclusion box, found {hits}"
prs.save(D)
