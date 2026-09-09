# -*- coding: utf-8 -*-
"""State the edition on L3's book slide, so page references cannot rot silently."""
import copy
from pptx import Presentation
from lxml import etree

A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
D="/Users/johnny/surfdrive/Teaching/RMS/RMS_2026/RMS_2026/RMS2627_3_WhyStatistics.pptx"
LINE="5th Global Edition, which all page references in these slides follow"

prs=Presentation(D)
s=list(prs.slides)[24]
body=next(sh for sh in s.shapes if sh.has_text_frame and "Do we all" in sh.text_frame.text)
tf=body.text_frame
for pa in tf.paragraphs:                      # idempotent
    if "Global Edition" in pa.text:
        pa._p.getparent().remove(pa._p)
paras=tf.paragraphs
anchor=next(i for i,p in enumerate(paras) if "Do we all have the book" in p.text)
model=paras[anchor]._p

q=copy.deepcopy(model)
for r in q.findall(A+"r")[1:]: q.remove(r)
run=q.find(A+"r")
for t in run.findall(A+"t"): run.remove(t)
etree.SubElement(run, A+"t").text=LINE
for rPr in run.findall(A+"rPr"): rPr.set("sz","1800")
pPr=q.find(A+"pPr")
if pPr is None:
    pPr=etree.Element(A+"pPr"); q.insert(0,pPr)
pPr.set("lvl","1")
paras[anchor]._p.addnext(q)
prs.save(D)
print(f"  slide 25: added sub-bullet {LINE!r}")
