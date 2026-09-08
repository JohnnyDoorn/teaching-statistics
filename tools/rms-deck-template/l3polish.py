# -*- coding: utf-8 -*-
"""Three fixes to slide 58: the purple key line, box placement, box text sizes."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt

D=sys.argv[1]
prs=Presentation(D)
s=list(prs.slides)[57]
log=[]

for sh in s.shapes:
    if not sh.has_text_frame: continue
    tf=sh.text_frame; txt=tf.text

    # 1. purple key line: it never loses, which is half its appeal
    if "not matter for your wallet" in txt:
        for pa in tf.paragraphs:
            if pa.text.strip().startswith("Purple:"):
                keep=pa.runs[0]
                keep.text="Purple: chasing a jackpot, while never losing"
                for r in pa.runs[1:]: r._r.getparent().remove(r._r)
                log.append("  key line -> 'Purple: chasing a jackpot, while never losing'")
        # 2. drop below the axis line so both curves stay fully visible
        old=(round(sh.left/914400,2), round(sh.top/914400,2),
             round(sh.width/914400,2), round(sh.height/914400,2))
        sh.left, sh.top, sh.width, sh.height = (Inches(0.92), Inches(5.80),
                                                Inches(11.50), Inches(1.55))
        tf.word_wrap = True
        # the empty spacer paragraph inherits 28pt from the master, which is
        # what pushed the last key line off the slide
        A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
        for pa in tf.paragraphs:
            if pa.text.strip(): continue
            end = pa._p.find(A+"endParaRPr")
            if end is None:
                from lxml import etree
                end = etree.SubElement(pa._p, A+"endParaRPr")
            end.set("sz", "1000")
            log.append("  blank spacer 28pt (inherited) -> 10pt")
        log.append(f"  conclusion box {old} -> (0.92, 5.8, 11.5, 1.55), full-width, clears both curves")

    # 3. match the two median boxes' explanation lines
    if txt.startswith("Median ≈"):
        for pa in list(tf.paragraphs)[1:]:
            for r in pa.runs:
                if r.font.size != Pt(16):
                    r.font.size=Pt(16)
                    log.append(f"  '{txt[:12].strip()}' explanation line -> 16pt")

prs.save(D)
print("\n".join(log))
