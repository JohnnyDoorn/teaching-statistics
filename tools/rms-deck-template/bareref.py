# -*- coding: utf-8 -*-
"""Bare page reference, placed where a core term is first introduced.

No explanatory sentence: the slide already says what the term is, so the
reference only needs to say where the book defines it.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

D="/Users/johnny/surfdrive/Teaching/RMS/RMS_2026/RMS_2026/RMS2627_4_VariabilityAndAssociation.pptx"
prs=Presentation(D); sl=list(prs.slides)

# 1. strip my prose from the standard-deviation reference
for sh in sl[7].shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Agresti, p. 91"):
        pa=sh.text_frame.paragraphs[0]
        pa.runs[0].text="(Agresti, p. 91)"
        for r in pa.runs[1:]: r._r.getparent().remove(r._r)
        print("  s8  -> '(Agresti, p. 91)'   [standard deviation first defined here]")

# 2. drop the reference off the film-still slide; it introduces no term
for sh in list(sl[38].shapes):
    if sh.has_text_frame and "Agresti, p. 137" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
        print("  s39 -> removed (title/image slide, no term introduced)")

# 3. put it where 'Contingency Table' is actually introduced
s41=sl[40]
for sh in list(s41.shapes):
    if sh.has_text_frame and "Agresti, p. 137" in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)
tf=s41.shapes.add_textbox(Inches(0.78), Inches(6.30), Inches(2.30), Inches(0.34)).text_frame
tf.word_wrap=False
r=tf.paragraphs[0].add_run(); r.text="(Agresti, p. 137)"
r.font.size=Pt(14); r.font.color.rgb=RGBColor(0,0,0)
print("  s41 -> '(Agresti, p. 137)'  [Contingency Table first introduced here]")
prs.save(D)
