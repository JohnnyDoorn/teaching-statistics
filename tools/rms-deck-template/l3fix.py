# -*- coding: utf-8 -*-
"""L3: interpret the two slot machines, and close the recap on why we do this."""
import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from lxml import etree

A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
D=sys.argv[1]
prs=Presentation(D)
sl=list(prs.slides)

# ---------- 1. new interpretation slide after 57 ----------
src=sl[56]                                   # slide 57, the two medians
box={}
for sh in src.shapes:
    if sh.has_text_frame and "Median ≈" in sh.text_frame.text:
        box["4.37" if "4.37" in sh.text_frame.text else "5.63"]=sh
assert len(box)==2, box.keys()

lay=next(l for l in prs.slide_masters[0].slide_layouts if l.name=="Title Only")
new=prs.slides.add_slide(lay)
new.shapes.title.text="Same mean, different game"

TEXT={"4.37":["Median 4.37, below the mean of 5",
              "Most sessions pay less than average. The mean is held up by "
              "the rare big win out in the long right tail.",
              "You never lose, but a typical night disappoints.",
              "Draws the player chasing the jackpot."],
      "5.63":["Median 5.63, above the mean of 5",
              "Most sessions pay more than average. The mean is dragged down "
              "by the rare heavy loss in the long left tail.",
              "A typical night is good, but you can walk away down.",
              "Draws the player who wants a reliable evening."]}
for key,(L,T) in (("4.37",(0.92,2.15)), ("5.63",(6.75,2.15))):
    el=copy.deepcopy(box[key]._element)
    new.shapes._spTree.append(el)
    shp=new.shapes[-1]
    shp.left, shp.top, shp.width, shp.height = Inches(L), Inches(T), Inches(5.66), Inches(2.75)
    tf=shp.text_frame; tf.word_wrap=True
    for extra in list(tf.paragraphs)[1:]: extra._p.getparent().remove(extra._p)
    p0=tf.paragraphs[0]
    for r in list(p0.runs): r._r.getparent().remove(r._r)
    for j,line in enumerate(TEXT[key]):
        para=p0 if j==0 else tf.add_paragraph()
        r=para.add_run(); r.text=line
        r.font.size=Pt(17 if j else 19); r.font.bold=(j==0)

tb=new.shapes.add_textbox(Inches(0.92), Inches(5.35), Inches(11.5), Inches(1.1)).text_frame
tb.word_wrap=True
r=tb.paragraphs[0].add_run()
r.text=("Both machines have mean = 5, so the mean alone cannot tell them apart. "
        "The median tells you what a typical night actually looks like.")
r.font.size=Pt(20)

lst=prs.slides._sldIdLst
mv=list(lst)[-1]; lst.remove(mv); lst.insert(57, mv)     # -> slide 58
print("  inserted 'Same mean, different game' as slide 58")

# ---------- 2. expand the recap's last point ----------
recap=next(s for s in prs.slides if s.shapes.title is not None
           and s.shapes.title.text.strip()=="Recap of Today")
body=next(sh for sh in recap.shapes if sh.has_text_frame and "statistics to" in sh.text_frame.text)
paras=body.text_frame.paragraphs
# name the two jobs on the bullets that already describe them, keeping the
# original wording and appending the label
for pa in paras:
    if pa.text.startswith("Get an overview of data") and "descriptive" not in pa.text:
        pa.runs[-1].text = pa.runs[-1].text + "  \u2192 descriptive statistics"
    if pa.text.startswith("Make statements about the whole population") and "inferential" not in pa.text:
        pa.runs[-1].text = pa.runs[-1].text + "  \u2192 inferential statistics"

# reclaim the blank spacers so the new line clears the arrows below
for pa in [p for p in paras if not p.text.strip()][:2]:
    pa._p.getparent().remove(pa._p)
paras=body.text_frame.paragraphs
anchor=next(i for i,p in enumerate(paras) if "Science quality" in p.text)
model=paras[1]._p                                        # an existing 22pt bullet
NEW=["do both well, or findings do not hold up: that is the replication crisis"]
for text in reversed(NEW):
    q=copy.deepcopy(model)
    for r in q.findall(A+"r")[1:]: q.remove(r)
    run=q.find(A+"r")
    for t in run.findall(A+"t"): run.remove(t)
    etree.SubElement(run, A+"t").text=text
    for rPr in run.findall(A+"rPr"): rPr.set("sz","1800")
    pPr=q.find(A+"pPr")
    if pPr is None:
        pPr=etree.Element(A+"pPr"); q.insert(0,pPr)
    pPr.set("lvl","1")
    paras[anchor]._p.addnext(q)
print("  recap: named descriptive/inferential, added the consequence line")
prs.save(D)
