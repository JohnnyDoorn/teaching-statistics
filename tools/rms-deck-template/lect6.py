# -*- coding: utf-8 -*-
"""Lecture 6: open on the roulette run, and close the loop in the recap."""
import copy, sys
from pptx import Presentation
from pptx.util import Pt
from lxml import etree

A="{http://schemas.openxmlformats.org/drawingml/2006/main}"
DECK=sys.argv[1]
prs=Presentation(DECK)

# ---- 1. reorder: 1, [4..10 roulette], 2, 3, 11.. -------------------------
lst = prs.slides._sldIdLst
ids = list(lst)
new = [ids[0]] + ids[3:10] + ids[1:3] + ids[10:]
assert len(new) == len(ids) and {id(x) for x in new} == {id(x) for x in ids}
for e in ids: lst.remove(e)
for e in new: lst.append(e)
print(f"  reordered: roulette now slides 2-8, depression/base-rate 9-10")

# ---- 2. expand the "what you condition on" recap point -------------------
NEW = ["our opening test was 99% accurate, yet only 4.7% of its positives were correct",
       "a rare condition lets the base rate dominate: report prevalence, not just "
       "sensitivity and specificity"]

recap = [s for s in prs.slides
         if s.shapes.title is not None and s.shapes.title.text.strip() == "Recap of Today"]
assert len(recap) == 1, f"expected 1 recap slide, found {len(recap)}"
body = next(sh for sh in recap[0].shapes
            if sh.has_text_frame and "conditional" in sh.text_frame.text)
paras = body.text_frame.paragraphs
anchor = next(i for i, p in enumerate(paras) if "what you condition on" in p.text)
model  = paras[1]._p                      # an existing lvl-1 sub-bullet, for styling

# reclaim a blank spacer above, so the new lines cannot reach the table below
for p in paras[:anchor]:
    if not p.text.strip():
        p._p.getparent().remove(p._p); break
paras  = body.text_frame.paragraphs
anchor = next(i for i, p in enumerate(paras) if "what you condition on" in p.text)

added = []
for text in reversed(NEW):
    p = copy.deepcopy(model)
    for r in p.findall(A+"r")[1:]: p.remove(r)          # keep one run
    run = p.find(A+"r")
    for t in run.findall(A+"t"): run.remove(t)
    t = etree.SubElement(run, A+"t"); t.text = text
    for rPr in run.findall(A+"rPr"): rPr.set("sz", "1600")
    paras[anchor]._p.addnext(p)
    added.append(p)
# drop one trailing blank so the block does not grow into the table
for p in list(body.text_frame.paragraphs)[::-1]:
    if not p.text.strip() and p._p not in added:
        p._p.getparent().remove(p._p); break
print(f"  recap: added {len(NEW)} sub-bullets under the P(A|B) point")

prs.save(DECK)
