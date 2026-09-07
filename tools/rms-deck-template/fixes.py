# -*- coding: utf-8 -*-
"""Housekeeping sweep: stale slide, typos, notation, and two hook placeholders."""
import copy, sys
from pptx import Presentation
from pptx.util import Pt

D = "/Users/johnny/surfdrive/Teaching/RMS/RMS_2026/RMS_2026/"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
log = []

def runs(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for pa in sh.text_frame.paragraphs:
            for r in pa.runs: yield r

# ---- 1. text replacements (exact, counted) ------------------------------
REPL = {"RMS2627_3_WhyStatistics.pptx":       [("=/=", "≠")],
        "RMS2627_6_ConditionalProbability.pptx": [("=/=", "≠")],
        "RMS2627_18_CategoricalAssociation.pptx": [("=/=", "≠")],
        "RMS2627_16_ComparingTwoGroups.pptx":  [("excercises", "exercises")],
        "RMS2627_20_NonparametricTests.pptx":  [("excercises", "exercises")]}
for fn, pairs in REPL.items():
    p = Presentation(D + fn); n = 0
    for i, s in enumerate(p.slides, 1):
        for r in runs(s):
            for old, new in pairs:
                if old in r.text:
                    r.text = r.text.replace(old, new); n += 1
    assert n, f"no match in {fn}"
    p.save(D + fn); log.append(f"  {fn:<42} {n} replacement(s) {pairs}")

# ---- 2. delete the stale strike slide -----------------------------------
fn = "RMS2627_16_ComparingTwoGroups.pptx"
p = Presentation(D + fn)
target = list(p.slides)[1]
assert "Strike" in " ".join(sh.text_frame.text for sh in target.shapes if sh.has_text_frame)
lst = p.slides._sldIdLst
sid = list(lst)[1]
p.part.drop_rel(sid.rId); lst.remove(sid)
p.save(D + fn); log.append(f"  {fn:<42} removed stale slide 2 (strike announcement)")

# ---- 3. hook placeholders ----------------------------------------------
HOOKS = {"RMS2627_20_NonparametricTests.pptx":
           ("Does this group really differ?",
            ["[ placeholder: scenario where the two groups look different, but the "
             "gap is driven by a single outlier ]",
             "[ then: what the t-test does with that value, and why a rank-based "
             "test does not ]"]),
         "RMS2627_24_MoreBayes.pptx":
           ("[ hook: opening example ]",
            ["[ placeholder: a result that looks convincing one way, and different "
             "once you ask a Bayesian question of it ]"])}
for fn, (title, bullets) in HOOKS.items():
    p = Presentation(D + fn)
    lay = next(l for l in p.slide_masters[0].slide_layouts if l.name == "Title and Content")
    s = p.slides.add_slide(lay)
    s.shapes.title.text = title
    body = next(ph for ph in s.placeholders if ph.placeholder_format.idx == 1)
    tf = body.text_frame
    for j, b in enumerate(bullets):
        para = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        r = para.add_run(); r.text = b; r.font.size = Pt(24)
    lst = p.slides._sldIdLst
    new = list(lst)[-1]; lst.remove(new); lst.insert(1, new)   # -> slide 2
    p.save(D + fn); log.append(f"  {fn:<42} inserted hook placeholder at slide 2")

print("\n".join(log))
