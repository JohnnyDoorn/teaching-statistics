# -*- coding: utf-8 -*-
"""Resolve old-template inheritance into explicit slide-level properties.

Slides pasted in from other decks carry placeholders whose geometry and bullet
formatting live in a layout we are about to delete.  Bake those values onto the
slide itself so the swap cannot change how they render.
"""
import copy, sys
from pptx import Presentation
from lxml import etree

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
SRC, DST = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

def lvl_pPr(shape, level):
    """Effective <a:lvlNpPr> for a placeholder level: layout lstStyle, else master."""
    ph = shape.placeholder_format.type
    layout = shape.part.slide_layout if hasattr(shape.part, "slide_layout") else None
    tag = f"{A}lvl{level+1}pPr"
    for src in (layout, ):
        if src is None: continue
        for lp in src.placeholders:
            if lp.placeholder_format.idx != shape.placeholder_format.idx: continue
            ls = lp._element.find(f"{P}txBody/{A}lstStyle")
            if ls is not None:
                el = ls.find(tag)
                if el is not None: return el
    master = layout.slide_master if layout is not None else None
    if master is None: return None
    styles = master._element.find(f"{P}txStyles")
    if styles is None: return None
    which = "bodyStyle" if str(ph).startswith(("BODY", "SUBTITLE", "OBJECT")) else "titleStyle"
    st = styles.find(f"{P}{which}")
    return None if st is None else st.find(tag)

baked = {"xfrm": 0, "bullets": 0, "slides": set()}
for n, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if not sh.is_placeholder: continue
        el = sh._element
        spPr = el.find(f"{P}spPr")
        # 1. geometry
        if spPr is not None and spPr.find(f"{A}xfrm") is None:
            try:
                L, T, Wd, H = sh.left, sh.top, sh.width, sh.height
            except Exception:
                L = None
            if None not in (L, T, Wd, H):
                x = etree.SubElement(spPr, f"{A}xfrm")
                etree.SubElement(x, f"{A}off", x=str(L), y=str(T))
                etree.SubElement(x, f"{A}ext", cx=str(Wd), cy=str(H))
                spPr.insert(0, x)
                baked["xfrm"] += 1; baked["slides"].add(n)
        # 2. bullet / indent formatting, per paragraph
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            p = para._p
            lvl = int(p.find(f"{A}pPr").get("lvl", 0)) if p.find(f"{A}pPr") is not None else 0
            src = lvl_pPr(sh, lvl)
            if src is None: continue
            pPr = p.find(f"{A}pPr")
            if pPr is None:
                pPr = etree.Element(f"{A}pPr"); p.insert(0, pPr)
            changed = False
            for attr in ("marL", "indent", "algn"):
                if pPr.get(attr) is None and src.get(attr) is not None:
                    pPr.set(attr, src.get(attr)); changed = True
            have_bu = any(pPr.find(f"{A}{t}") is not None
                          for t in ("buNone", "buChar", "buAutoNum"))
            if not have_bu:
                for t in ("buClr", "buSzPct", "buSzPts", "buFont", "buChar", "buNone", "buAutoNum"):
                    e = src.find(f"{A}{t}")
                    if e is not None: pPr.append(copy.deepcopy(e)); changed = True
            if changed: baked["bullets"] += 1; baked["slides"].add(n)

prs.save(DST)
print(f"  baked {baked['xfrm']} xfrm + {baked['bullets']} paragraph styles "
      f"on slides {sorted(baked['slides'])}")
