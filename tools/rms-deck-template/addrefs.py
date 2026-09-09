# -*- coding: utf-8 -*-
"""Two extra Agresti references for L4, in the deck's own footnote style (14pt)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

D="/Users/johnny/surfdrive/Teaching/RMS/RMS_2026/RMS_2026/RMS2627_4_VariabilityAndAssociation.pptx"
ADD={8:  ((0.92, 6.40, 7.60, 0.42),
          "Agresti p. 91 works through the standard deviation step by step"),
     39: ((0.34, 6.62, 3.90, 0.55),
          "Agresti p. 137 on association between two categorical variables")}
prs=Presentation(D)
for idx,((L,T,W,H), text) in sorted(ADD.items()):
    s=list(prs.slides)[idx-1]
    for sh in list(s.shapes):                       # idempotent
        if sh.has_text_frame and text[:18] in sh.text_frame.text:
            sh._element.getparent().remove(sh._element)
    tf=s.shapes.add_textbox(Inches(L), Inches(T), Inches(W), Inches(H)).text_frame
    tf.word_wrap=True
    r=tf.paragraphs[0].add_run(); r.text=text
    r.font.size=Pt(14); r.font.color.rgb=RGBColor(0,0,0)
    print(f"   slide {idx}: {text}")
prs.save(D)
