# -*- coding: utf-8 -*-
"""One citation form: 'Agresti, p. N'.

Left alone where the sentence already names the authors immediately before the
reference -- 'Agresti & Franklin: "..." (p. 26)' is correct as it stands, and
repeating the name there would be redundant.
"""
from pptx import Presentation

D = "/Users/johnny/surfdrive/Teaching/RMS/RMS_2026/RMS_2026/"
FIXES = [
 ("RMS2627_3_WhyStatistics.pptx",            [32],       "(p. 58)",   "(Agresti, p. 58)"),
 ("RMS2627_4_VariabilityAndAssociation.pptx",[8],        "Agresti p. 91",  "Agresti, p. 91"),
 ("RMS2627_4_VariabilityAndAssociation.pptx",[31],       "Agresti p. 99",  "Agresti, p. 99"),
 ("RMS2627_4_VariabilityAndAssociation.pptx",[39],       "Agresti p. 137", "Agresti, p. 137"),
 ("RMS2627_4_VariabilityAndAssociation.pptx",[50],       "(Agresti p. 154)","(Agresti, p. 154)"),
 ("RMS2627_20_NonparametricTests.pptx",      [27],       "(p. 464,",  "(Agresti, p. 464,"),
 ("RMS2627_20_NonparametricTests.pptx",      [61],       "on p. 464", "on Agresti, p. 464"),
 ("RMS2627_20_NonparametricTests.pptx",      [83],       "(see p. 464)","(see Agresti, p. 464)"),
]

def replace(slide, old, new):
    n = 0
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for pa in sh.text_frame.paragraphs:
            runs = pa.runs
            if not runs: continue
            joined = "".join(r.text for r in runs)
            if old not in joined: continue
            for r in runs:
                if old in r.text:
                    r.text = r.text.replace(old, new); n += 1; break
            else:
                runs[0].text = joined.replace(old, new)
                for r in runs[1:]: r._r.getparent().remove(r._r)
                n += 1
    return n

by_deck = {}
for deck, slides, old, new in FIXES:
    by_deck.setdefault(deck, []).append((slides, old, new))
for deck, jobs in by_deck.items():
    prs = Presentation(D + deck); sl = list(prs.slides)
    for slides, old, new in jobs:
        hits = sum(replace(sl[i-1], old, new) for i in slides)
        assert hits == len(slides), f"{deck} {old!r}: expected {len(slides)}, got {hits}"
        print(f"  {deck.split('_')[1]:<3} s{str(slides):<8} {old!r} -> {new!r}")
    prs.save(D + deck)
