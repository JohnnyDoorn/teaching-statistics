# -*- coding: utf-8 -*-
"""Make the 3.17 answer readable on its own.

The book's appendix gives only '1.(c); 2.(a); 3.(d); 4.(b)', and the four
candidate r values live inside the figure rather than the text, so the slide has
to spell out the pairing or it means nothing away from the book.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

D="/Users/johnny/surfdrive/Teaching/RMS/RMS_2026/RMS_2026/RMS2627_4_VariabilityAndAssociation.pptx"
LINES=[
 "Each scatterplot gets one of the four correlation coefficients:",
 "",
 "Scatterplot 1  →  r = −0.9   (option c)    strong negative: points hug a downward line",
 "Scatterplot 2  →  r = −0.5   (option a)    moderate negative: downward, but scattered",
 "Scatterplot 3  →  r = 0       (option d)    no linear association",
 "Scatterplot 4  →  r = 0.6    (option b)    moderate positive: upward, fairly scattered",
 "",
 "The sign tells you the direction, the size tells you how tightly the points "
 "follow a straight line.",
]
prs=Presentation(D)
target=next(s for s in prs.slides
            if s.shapes.title is not None and s.shapes.title.text.strip().startswith("3.17"))
for sh in list(target.shapes):
    if sh.has_text_frame and sh is not target.shapes.title and sh.text_frame.text.strip():
        sh._element.getparent().remove(sh._element)
tf=target.shapes.add_textbox(Inches(0.92), Inches(2.05), Inches(11.4), Inches(4.0)).text_frame
tf.word_wrap=True
for j,l in enumerate(LINES):
    pa=tf.paragraphs[0] if j==0 else tf.add_paragraph()
    if not l.strip():
        pa.add_run().text=""
        for r in pa.runs: r.font.size=Pt(10)
        continue
    r=pa.add_run(); r.text=l
    r.font.size=Pt(19 if j==0 or j==len(LINES)-1 else 18)
    r.font.color.rgb=RGBColor(0,0,0)
    r.font.bold=(j==0)
tf2=target.shapes.add_textbox(Inches(0.92), Inches(6.25), Inches(11.4), Inches(0.5)).text_frame
rr=tf2.paragraphs[0].add_run()
rr.text=("Letter matching from Agresti's answer section; the four r values read off the "
         "figure and checked by JvD")
rr.font.size=Pt(14)
prs.save(D)
print("  3.17 answer slide rewritten")
