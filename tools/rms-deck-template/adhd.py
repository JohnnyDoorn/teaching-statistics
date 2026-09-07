# -*- coding: utf-8 -*-
"""Lecture 6: one disorder throughout (ADHD), with a defensible base rate."""
import sys
from pptx import Presentation

DECK = sys.argv[1]
prs  = Presentation(DECK)
S    = lambda n: list(prs.slides)[n - 1]

# (slide, exact old run text, new run text)
EDITS = [
 # -- opening puzzle: depression @ prev .001 -> ADHD @ prev .05, noisy test --
 (9, "Probability of depression is 0.001 (prevalence/ base rate)",
     "Probability of ADHD in children is 0.05 (prevalence/ base rate)"),
 (9, "Probability that the diagnostic test indicates that you have a depression, "
     "when you really have a depression is 0.99",
     "Probability that the diagnostic test indicates that you have ADHD, "
     "when you really have ADHD is 0.90"),
 (9, " have a depression, when indeed you don’t have a depression is 0.98",
     " have ADHD, when indeed you don’t have ADHD is 0.75"),
 (9, "What is the probability that you have a depression, when the test indicates "
     "that you have a depression?",
     "What is the probability that you have ADHD, when the test indicates "
     "that you have ADHD?"),
 (9, ")=0.0472!", ")=0.1593!"),
 # -- example list: autism/CARS -> ADHD (no invented instrument or cut-off) --
 (13, "Having autism ", "Having ADHD "),
 (13, " CARS score > 20", " the rating scale score is above the cut-off"),
 # -- recap: numbers must follow the new opening --
 (70, "our opening test was 99% accurate, yet only 4.7% of its positives were correct",
      "our opening test caught 90% of real cases, yet only 16% of its positives were correct"),
 (70, "a rare condition lets the base rate dominate: report prevalence, not just "
      "sensitivity and specificity",
      "a low base rate plus an imperfect test makes most positives false: "
      "prevalence matters too"),
]

done = []
for n, old, new in EDITS:
    hit = 0
    for sh in S(n).shapes:
        if not sh.has_text_frame: continue
        for pa in sh.text_frame.paragraphs:
            for r in pa.runs:
                if r.text == old:
                    r.text = new; hit += 1
    if hit != 1:
        raise SystemExit(f"ABORT slide {n}: matched {hit} runs for {old[:48]!r}")
    done.append((n, old[:44], new[:44]))

for n, o, nw in done:
    print(f"  s{n:>3}  {o!r}\n        -> {nw!r}")
prs.save(DECK)
print(f"\n{len(done)} run replacements, all uniquely matched")
