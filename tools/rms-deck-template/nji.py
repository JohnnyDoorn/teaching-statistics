# -*- coding: utf-8 -*-
"""Source note for the ADHD base rate on the opening puzzle slide."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

DECK = sys.argv[1]
URL  = "https://www.nji.nl/databanken/cijfers/cijfers-over-adhd"
LEAD = "ADHD prevalence: 5.3% worldwide, 4.7% in NL "
CITE = "(Nederlands Jeugdinstituut)"

prs   = Presentation(DECK)
slide = list(prs.slides)[8]                       # opening ADHD puzzle
for sh in list(slide.shapes):                     # idempotent
    if sh.has_text_frame and CITE in sh.text_frame.text:
        sh._element.getparent().remove(sh._element)

tf = slide.shapes.add_textbox(Inches(0.92), Inches(6.82), Inches(8.0), Inches(0.28)).text_frame
tf.word_wrap = False
p = tf.paragraphs[0]
r1 = p.add_run(); r1.text = LEAD
r1.font.size, r1.font.name = Pt(10), "Arial"
r1.font.color.rgb = RGBColor(0, 0, 0)
r2 = p.add_run(); r2.text = CITE
r2.font.size, r2.font.name = Pt(10), "Arial"
r2.hyperlink.address = URL
prs.save(DECK)
print(f"  slide 9: {LEAD}{CITE}  -> {URL}")
